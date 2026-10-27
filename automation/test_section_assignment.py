from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.by import By
import time
from pathlib import Path
from openpyxl import load_workbook


# ==========================================================
# READ EXTERNAL DATA FROM EXCEL
# ==========================================================

file_path = Path(__file__).resolve().parent / "test_data.xlsx"
workbook = load_workbook(file_path)

sheet = workbook["Sheet1"]


# ==========================================================
# READ DATA USING EXCEL COLUMN HEADERS
# ==========================================================

headers = [cell.value for cell in sheet[1]]

data = {}

for index, header in enumerate(headers, start=1):

    data[header] = sheet.cell(
        row=2,
        column=index
    ).value


# ==========================================================
# TEST DATA FROM EXCEL
# ==========================================================

branch_value = data["Branch"]
department_value = data["Department"]
program_value = data["Program"]
semester_value = data["Semester"]
semester_label_value = data["Semester Label"]

count_1_value = str(data["Count 1"])
count_2_value = str(data["Count 2"])

sort_by_value = data["Sort By"]
order_by_value = data["Order By"]


# ==========================================================
# TEST
# ==========================================================

def test_section_assignment(driver):

    wait = WebDriverWait(driver, 15)


    # ==========================================================
    # OPEN STUDENT INFO
    # ==========================================================

    student_info_menu = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//div//span[normalize-space()='Student Info']"
            )
        )
    )

    student_info_menu.click()


    # ==========================================================
    # OPEN SECTION ASSIGNMENT
    # ==========================================================

    section_assignment = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//li[contains(.,'Section Assignment')])[2]"
            )
        )
    )

    section_assignment.click()


    # ==========================================================
    # BRANCH DROPDOWN
    # ==========================================================

    branch = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[1]"
            )
        )
    )

    branch.click()

    branch_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{branch_value}']"
            )
        )
    )

    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});",
        branch_option
    )

    driver.execute_script(
        "arguments[0].click();",
        branch_option
    )


    # ==========================================================
    # DEPARTMENT DROPDOWN
    # ==========================================================

    department = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[2]"
            )
        )
    )

    department.click()

    department_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{department_value}']"
            )
        )
    )

    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});",
        department_option
    )

    department_option.click()


    # ==========================================================
    # PROGRAM DROPDOWN
    # ==========================================================

    program = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[3]"
            )
        )
    )

    program.click()

    program_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{program_value}']"
            )
        )
    )

    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});",
        program_option
    )

    program_option.click()


    # ==========================================================
    # SEMESTER DROPDOWN
    # ==========================================================

    semester = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[4]"
            )
        )
    )

    semester.click()

    semester_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{semester_value}']"
            )
        )
    )

    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});",
        semester_option
    )

    semester_option.click()


    # ==========================================================
    # SEMESTER LABEL DROPDOWN
    # ==========================================================

    semester_label = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[5]"
            )
        )
    )

    semester_label.click()

    semester_label_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{semester_label_value}']"
            )
        )
    )

    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});",
        semester_label_option
    )

    semester_label_option.click()


    # ==========================================================
    # SEARCH
    # ==========================================================

    search_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.ID,
                "fetchStudents"
            )
        )
    )

    search_button.click()


    # ==========================================================
    # SCROLL DOWN
    # ==========================================================

    driver.execute_script(
        "window.scrollTo(0, document.body.scrollHeight);"
    )


    # ==========================================================
    # ASSIGN SECTION BUTTON
    # ==========================================================

    time.sleep(2)

    assign_section_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.ID,
                "openAssignedSectionBtn"
            )
        )
    )

    assign_section_button.click()


    # ==========================================================
    # SECTION A CHECKBOX
    # ==========================================================

    section_a_checkbox = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//input[@class='section-config-toggle'])[1]"
            )
        )
    )

    if not section_a_checkbox.is_selected():

        section_a_checkbox.click()

    section_a_checkbox = wait.until(
        EC.presence_of_element_located(
            (
                By.XPATH,
                "(//input[@class='section-config-toggle'])[1]"
            )
        )
    )

    assert section_a_checkbox.is_selected()


    # ==========================================================
    # SECTION B CHECKBOX
    # ==========================================================

    section_b_checkbox = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//input[@class='section-config-toggle'])[2]"
            )
        )
    )

    if not section_b_checkbox.is_selected():

        section_b_checkbox.click()

    section_b_checkbox = wait.until(
        EC.presence_of_element_located(
            (
                By.XPATH,
                "(//input[@class='section-config-toggle'])[2]"
            )
        )
    )

    assert section_b_checkbox.is_selected()


    # ==========================================================
    # COUNT FIELD 1
    # ==========================================================

    count_field_1 = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//input[contains(@class,'section-count')])[1]"
            )
        )
    )

    count_field_1.clear()

    count_field_1.send_keys(count_1_value)

    assert count_field_1.get_attribute(
        "value"
    ) == count_1_value


    # ==========================================================
    # COUNT FIELD 2
    # ==========================================================

    count_field_2 = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//input[contains(@class,'section-count')])[2]"
            )
        )
    )

    count_field_2.clear()

    count_field_2.send_keys(count_2_value)

    assert count_field_2.get_attribute(
        "value"
    ) == count_2_value


    # ==========================================================
    # SORT BY
    # ==========================================================

    sort_by_dropdown = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[8]"
            )
        )
    )

    sort_by_dropdown.click()

    sort_by_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{sort_by_value}']"
            )
        )
    )

    sort_by_option.click()


    # ==========================================================
    # ORDER BY
    # ==========================================================

    order_by_dropdown = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "(//div[@class='nice-select primary_select form-control'])[9]"
            )
        )
    )

    order_by_dropdown.click()

    order_by_option = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                f"//li[normalize-space()='{order_by_value}']"
            )
        )
    )

    order_by_option.click()


    # ==========================================================
    # PREVIEW
    # ==========================================================

    preview_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.ID,
                "previewBtn"
            )
        )
    )

    preview_button.click()


    # ==========================================================
    # SAVE
    # ==========================================================

    save_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.ID,
                "saveBtn"
            )
        )
    )

    save_button.click()


    # ==========================================================
    # CONFIRM POPUP
    # ==========================================================

    confirm_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//button[contains(@class,'swal2-confirm')]"
            )
        )
    )

    confirm_button.click()

    time.sleep(2)