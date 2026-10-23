
# from openpyxl import Workbook
# from openpyxl.styles import Font, Alignment
# from openpyxl.utils import get_column_letter
# from datetime import datetime
# import os


# # ==========================================================
# # DOWNLOAD FOLDER
# # ==========================================================

# DOWNLOAD_FOLDER = os.path.join(
#     os.path.expanduser("~"),
#     "Downloads"
# )

# os.makedirs(
#     DOWNLOAD_FOLDER,
#     exist_ok=True
# )


# # ==========================================================
# # REPORT FILE
# # ==========================================================

# timestamp = datetime.now().strftime(
#     "%Y%m%d_%H%M%S"
# )

# REPORT_FILE = os.path.join(
#     DOWNLOAD_FOLDER,
#     f"test_execution_report_{timestamp}.xlsx"
# )


# def create_api_report(api_results):

#     workbook = Workbook()

#     # ======================================================
#     # API SUMMARY SHEET
#     # ======================================================

#     summary = workbook.active
#     summary.title = "Summary"

#     total = len(api_results)

#     passed = sum(
#         1 for result in api_results
#         if result.get("Status") == "PASS"
#     )

#     failed = sum(
#         1 for result in api_results
#         if result.get("Status") == "FAIL"
#     )

#     pass_percentage = (
#         (passed / total) * 100
#         if total > 0 else 0
#     )

#     summary["A1"] = "API Test Execution Summary"
#     summary["A1"].font = Font(bold=True, size=14)
#     summary.merge_cells("A1:B1")
#     summary["A1"].alignment = Alignment(horizontal="center")

#     summary.append(["Metric", "Value"])
#     summary.append(["Total API Tests", total])
#     summary.append(["Passed", passed])
#     summary.append(["Failed", failed])
#     summary.append(["Pass Percentage", f"{pass_percentage:.2f}%"])
#     summary.append([
#         "Execution Date",
#         datetime.now().strftime("%d-%m-%Y")
#     ])
#     summary.append([
#         "Execution Time",
#         datetime.now().strftime("%H:%M:%S")
#     ])

#     if total == 0:
#         report_status = "NO DATA"
#     elif failed > 0:
#         report_status = "FAILED"
#     else:
#         report_status = "PASSED"

#     summary.append(["Overall Status", report_status])
#     summary.append(["Report File", os.path.basename(REPORT_FILE)])

#     for cell in summary[2]:
#         cell.font = Font(bold=True)
#         cell.alignment = Alignment(horizontal="center")

#     for row in summary.iter_rows(min_row=3):
#         for cell in row:
#             cell.alignment = Alignment(
#                 vertical="center",
#                 wrap_text=True
#             )

#     # ======================================================
#     # API TEST RESULTS SHEET
#     # ======================================================

#     api_sheet = workbook.create_sheet("API Test Results")

#     api_sheet.append([
#         "Sr. No.",
#         "API Name",
#         "Staff Number",
#         "Test Case ID",
#         "Test Scenario",
#         "Method",
#         "Endpoint",
#         "Test Data / Payload",
#         "Status Code",
#         "Expected Result",
#         "Actual Result",
#         "Pass/Fail",
#         "Response Time (ms)",
#         "Error Details",
#         "Execution Date & Time"
#     ])

#     for index, result in enumerate(api_results, start=1):

#         api_sheet.append([
#             index,
#             result.get("API Name", ""),
#             result.get("Test Case ID", ""),
#             result.get("Test Scenario", ""),
#             result.get("Method", ""),
#             result.get("Endpoint", ""),
#             str(result.get("Test Data / Payload", "")),
#             result.get("Status Code", ""),
#             result.get("Expected Result", ""),
#             result.get("Actual Result", ""),
#             result.get("Status", ""),
#             result.get("Response Time (ms)", ""),
#             result.get("Error", ""),
#             result.get(
#                 "Execution Date & Time",
#                 datetime.now().strftime("%d-%m-%Y %H:%M:%S")
#             )
#         ])

#     for cell in api_sheet[1]:
#         cell.font = Font(bold=True)
#         cell.alignment = Alignment(
#             horizontal="center",
#             vertical="center",
#             wrap_text=True
#         )

#     for row in api_sheet.iter_rows(min_row=2):
#         for cell in row:
#             cell.alignment = Alignment(
#                 vertical="top",
#                 wrap_text=True
#             )

#     api_sheet.freeze_panes = "A2"
#     api_sheet.auto_filter.ref = api_sheet.dimensions
#     api_sheet.row_dimensions[1].height = 35
#     summary.freeze_panes = "A3"

#     # ======================================================
#     # AUTO COLUMN WIDTH
#     # ======================================================

#     for sheet in workbook.worksheets:
#         for column in sheet.columns:

#             max_length = 0
#             column_letter = get_column_letter(column[0].column)

#             for cell in column:
#                 if cell.value is not None:
#                     max_length = max(
#                         max_length,
#                         len(str(cell.value))
#                     )

#             sheet.column_dimensions[column_letter].width = min(
#                 max_length + 3,
#                 50
#             )

#     # ======================================================
#     # SAVE API REPORT
#     # ======================================================

#     workbook.save(REPORT_FILE)

#     print("\nAPI REPORT GENERATED SUCCESSFULLY")
#     print("Total API Tests:", total)
#     print("Passed:", passed)
#     print("Failed:", failed)
#     print("Report Status:", report_status)
#     print("Report Path:", os.path.abspath(REPORT_FILE))

# # ==========================================================
# # CREATE REPORT
# # ==========================================================


# def create_report(results, api_results=None, api_only=False):

#     # If this is an API-only report, create the API report
#     if api_only:
#         create_api_report(api_results or [])
#         return

#     # Existing UI report code continues below
#     workbook = Workbook()

#     # ======================================================
#     # SUMMARY SHEET
#     # ======================================================

#     summary = workbook.active
#     summary.title = "Summary"

#     # ======================================================
#     # GET MODULES DYNAMICALLY
#     # ======================================================

#     modules = []

#     for result in results:

#         module = result.get(
#             "Module",
#             result.get(
#                 "Test Case",
#                 "Unknown"
#             )
#         )

#         if module not in modules:

#             modules.append(module)

#     # ======================================================
#     # SUMMARY TITLE
#     # ======================================================

#     summary["A1"] = "Test Execution Summary"

#     summary["A1"].font = Font(
#         bold=True,
#         size=14
#     )

#     summary.merge_cells(
#         "A1:I1"
#     )

#     summary["A1"].alignment = Alignment(
#         horizontal="center"
#     )

#     # ======================================================
#     # SUMMARY HEADER
#     # ======================================================

#     summary.append([])

#     summary.append([
#         "Module",
#         "Total Test Cases",
#         "Passed",
#         "Failed",
#         "First Attempt Failures",
#         "Retry Attempts",
#         "Passed After Retry",
#         "Total Attempts",
#         "Pass %"
#     ])

#     # ======================================================
#     # MODULE-WISE SUMMARY
#     # ======================================================

#     for module in modules:

#         module_results = [
#             result
#             for result in results
#             if result.get(
#                 "Module",
#                 result.get(
#                     "Test Case",
#                     "Unknown"
#                 )
#             ) == module
#         ]

#         # --------------------------------------------------
#         # UNIQUE TEST DATA
#         # --------------------------------------------------

#         test_cases = []

#         for result in module_results:

#             test_data = result.get(
#                 "Test Data",
#                 ""
#             )

#             if test_data not in test_cases:

#                 test_cases.append(
#                     test_data
#                 )

#         total = len(test_cases)

#         # --------------------------------------------------
#         # FINAL STATUS
#         # --------------------------------------------------

#         final_status = {}

#         for result in module_results:

#             test_data = result.get(
#                 "Test Data",
#                 ""
#             )

#             status = result.get(
#                 "Final Status",
#                 ""
#             )

#             if status == "PASS":

#                 final_status[
#                     test_data
#                 ] = "PASS"

#             elif status == "FAIL":

#                 final_status[
#                     test_data
#                 ] = "FAIL"

#         # --------------------------------------------------
#         # PASSED
#         # --------------------------------------------------

#         passed = sum(
#             1
#             for status in final_status.values()
#             if status == "PASS"
#         )

#         # --------------------------------------------------
#         # FAILED
#         # --------------------------------------------------

#         failed = sum(
#             1
#             for status in final_status.values()
#             if status == "FAIL"
#         )

#         # --------------------------------------------------
#         # FIRST ATTEMPT FAILURES
#         # --------------------------------------------------

#         first_attempt_failures = sum(
#             1
#             for result in module_results
#             if result.get("Attempt") == 1
#             and result.get("Status") == "FAIL"
#         )

#         # --------------------------------------------------
#         # RETRY ATTEMPTS
#         # --------------------------------------------------

#         retry_attempts = sum(
#             1
#             for result in module_results
#             if result.get("Final Status") == "RETRY"
#         )

#         # --------------------------------------------------
#         # PASSED AFTER RETRY
#         # --------------------------------------------------

#         passed_after_retry = 0

#         for test_data in test_cases:

#             if final_status.get(
#                 test_data
#             ) == "PASS":

#                 retry_pass = any(
#                     result.get(
#                         "Test Data"
#                     ) == test_data
#                     and result.get(
#                         "Attempt",
#                         1
#                     ) > 1
#                     and result.get(
#                         "Status"
#                     ) == "PASS"

#                     for result in module_results
#                 )

#                 if retry_pass:

#                     passed_after_retry += 1

#         # --------------------------------------------------
#         # TOTAL ATTEMPTS
#         # --------------------------------------------------

#         total_attempts = len(
#             module_results
#         )

#         # --------------------------------------------------
#         # PASS %
#         # --------------------------------------------------

#         if total > 0:

#             pass_percentage = (
#                 passed / total
#             ) * 100

#         else:

#             pass_percentage = 0

#         # --------------------------------------------------
#         # ADD SUMMARY ROW
#         # --------------------------------------------------

#         summary.append([
#             module,
#             total,
#             passed,
#             failed,
#             first_attempt_failures,
#             retry_attempts,
#             passed_after_retry,
#             total_attempts,
#             f"{pass_percentage:.2f}%"
#         ])

#     # ======================================================
#     # SUMMARY HEADER STYLE
#     # ======================================================

#     for cell in summary[3]:

#         cell.font = Font(
#             bold=True
#         )

#         cell.alignment = Alignment(
#             horizontal="center",
#             vertical="center"
#         )

#     # ======================================================
#     # SUMMARY DATA ALIGNMENT
#     # ======================================================

#     for row in summary.iter_rows(
#         min_row=4
#     ):

#         for cell in row:

#             cell.alignment = Alignment(
#                 horizontal="center",
#                 vertical="center",
#                 wrap_text=True
#             )

#     # ======================================================
#     # EXECUTION INFORMATION
#     # ======================================================

#     execution_row = (
#         summary.max_row + 2
#     )

#     summary.cell(
#         execution_row,
#         1
#     ).value = "Execution Information"

#     summary.cell(
#         execution_row,
#         1
#     ).font = Font(
#         bold=True,
#         size=12
#     )

#     summary.cell(
#         execution_row + 1,
#         1
#     ).value = "Execution Date"

#     summary.cell(
#         execution_row + 1,
#         2
#     ).value = datetime.now().strftime(
#         "%d-%m-%Y"
#     )

#     summary.cell(
#         execution_row + 2,
#         1
#     ).value = "Execution Time"

#     summary.cell(
#         execution_row + 2,
#         2
#     ).value = datetime.now().strftime(
#         "%H:%M:%S"
#     )

#     # ======================================================
#     # OVERALL REPORT STATUS
#     # ======================================================

#     all_final_statuses = list(
#         result
#         for result in results
#         if result.get(
#             "Final Status"
#         ) in [
#             "PASS",
#             "FAIL"
#         ]
#     )

#     overall_failed = any(
#         result.get(
#             "Final Status"
#         ) == "FAIL"

#         for result in all_final_statuses
#     )

#     if not results:

#         report_status = "NO DATA"

#     elif overall_failed:

#         report_status = "FAILED"

#     else:

#         report_status = "PASSED"

#     summary.cell(
#         execution_row + 3,
#         1
#     ).value = "Overall Status"

#     summary.cell(
#         execution_row + 3,
#         2
#     ).value = report_status

#     summary.cell(
#         execution_row + 4,
#         1
#     ).value = "Report File"

#     summary.cell(
#         execution_row + 4,
#         2
#     ).value = os.path.basename(
#         REPORT_FILE
#     )

#     for row in range(
#         execution_row + 1,
#         execution_row + 5
#     ):

#         summary.cell(
#             row,
#             1
#         ).font = Font(
#             bold=True
#         )

#     # ======================================================
#     # TEST RESULTS SHEET
#     # ======================================================

#     details = workbook.create_sheet(
#         "Test Results"
#     )

#     details.append([
#         "Module",
#         "Test Case",
#         "Test Data",
#         "Attempt",
#         "Status",
#         "Error",
#         "Final Status"
#     ])

#     # ======================================================
#     # ADD RESULTS
#     # ======================================================

#     for result in results:

#         module = result.get(
#             "Module",
#             result.get(
#                 "Test Case",
#                 "Unknown"
#             )
#         )

#         details.append([
#             module,
#             result.get(
#                 "Test Case",
#                 ""
#             ),
#             result.get(
#                 "Test Data",
#                 ""
#             ),
#             result.get(
#                 "Attempt",
#                 ""
#             ),
#             result.get(
#                 "Status",
#                 ""
#             ),
#             result.get(
#                 "Error",
#                 ""
#             ),
#             result.get(
#                 "Final Status",
#                 ""
#             )
#         ])

#     # ======================================================
#     # DETAILS HEADER
#     # ======================================================

#     for cell in details[1]:

#         cell.font = Font(
#             bold=True
#         )

#         cell.alignment = Alignment(
#             horizontal="center",
#             vertical="center"
#         )

#     # ======================================================
#     # DETAILS ALIGNMENT
#     # ======================================================

#     for row in details.iter_rows(
#         min_row=2
#     ):

#         for cell in row:

#             cell.alignment = Alignment(
#                 vertical="top",
#                 wrap_text=True
#             )

#     # ======================================================
#     # FILTER
#     # ======================================================

#     if details.max_row > 1:

#         details.auto_filter.ref = (
#             details.dimensions
#         )

#     # ======================================================
#     # FREEZE PANES
#     # ======================================================

#     summary.freeze_panes = "A4"

#     details.freeze_panes = "A2"

#     # ======================================================
#     # AUTO COLUMN WIDTH
#     # ======================================================

#     for sheet in workbook.worksheets:

#         for column in sheet.columns:

#             max_length = 0

#             column_letter = (
#                 get_column_letter(
#                     column[0].column
#                 )
#             )

#             for cell in column:

#                 if cell.value is not None:

#                     length = len(
#                         str(cell.value)
#                     )

#                     if length > max_length:

#                         max_length = length

#             sheet.column_dimensions[
#                 column_letter
#             ].width = min(
#                 max_length + 3,
#                 50
#             )

#     # ======================================================
#     # SAVE REPORT
#     # ======================================================

#     workbook.save(
#         REPORT_FILE
#     )
    

#     # ======================================================
#     # CONSOLE OUTPUT
#     # ======================================================

#     print(
#         "\n===================================="
#     )

#     print(
#         "REPORT GENERATED SUCCESSFULLY"
#     )

#     print(
#         "===================================="
#     )

#     print(
#         f"Total Results : {len(results)}"
#     )

#     print(
#         f"Report Status : {report_status}"
#     )

#     print(
#         f"Report Path   : "
#         f"{os.path.abspath(REPORT_FILE)}"
#     )

#     print(
#         "===================================="
#     )




# from openpyxl import Workbook
# from openpyxl.styles import Font, Alignment
# from openpyxl.utils import get_column_letter
# from datetime import datetime
# import os


# # ==========================================================
# # DOWNLOAD FOLDER
# # ==========================================================

# DOWNLOAD_FOLDER = os.path.join(
#     os.path.expanduser("~"),
#     "Downloads"
# )

# os.makedirs(DOWNLOAD_FOLDER, exist_ok=True)


# # ==========================================================
# # REPORT FILE
# # ==========================================================

# timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

# REPORT_FILE = os.path.join(
#     DOWNLOAD_FOLDER,
#     f"test_execution_report_{timestamp}.xlsx"
# )


# # ==========================================================
# # COMMON: AUTO COLUMN WIDTH
# # ==========================================================

# def set_column_widths(workbook):

#     for sheet in workbook.worksheets:
#         for column in sheet.columns:

#             max_length = 0
#             column_letter = get_column_letter(column[0].column)

#             for cell in column:
#                 if cell.value is not None:
#                     max_length = max(
#                         max_length,
#                         len(str(cell.value))
#                     )

#             sheet.column_dimensions[column_letter].width = min(
#                 max_length + 3,
#                 50
#             )


# # ==========================================================
# # API REPORT
# # ==========================================================

# def create_api_report(api_results):

#     workbook = Workbook()

#     # ======================================================
#     # API SUMMARY SHEET
#     # ======================================================

#     summary = workbook.active
#     summary.title = "Summary"

#     total = len(api_results)

#     passed = sum(
#         1 for result in api_results
#         if result.get("Status") == "PASS"
#     )

#     failed = sum(
#         1 for result in api_results
#         if result.get("Status") == "FAIL"
#     )

#     pass_percentage = (
#         (passed / total) * 100
#         if total > 0 else 0
#     )

#     summary["A1"] = "API Test Execution Summary"
#     summary["A1"].font = Font(bold=True, size=14)
#     summary.merge_cells("A1:B1")
#     summary["A1"].alignment = Alignment(horizontal="center")

#     summary.append(["Metric", "Value"])
#     summary.append(["Total API Tests", total])
#     summary.append(["Passed", passed])
#     summary.append(["Failed", failed])
#     summary.append(["Pass Percentage", f"{pass_percentage:.2f}%"])
#     summary.append([
#         "Execution Date",
#         datetime.now().strftime("%d-%m-%Y")
#     ])
#     summary.append([
#         "Execution Time",
#         datetime.now().strftime("%H:%M:%S")
#     ])

#     if total == 0:
#         report_status = "NO DATA"
#     elif failed > 0:
#         report_status = "FAILED"
#     else:
#         report_status = "PASSED"

#     summary.append(["Overall Status", report_status])
#     summary.append(["Report File", os.path.basename(REPORT_FILE)])

#     for cell in summary[2]:
#         cell.font = Font(bold=True)
#         cell.alignment = Alignment(horizontal="center")

#     for row in summary.iter_rows(min_row=3):
#         for cell in row:
#             cell.alignment = Alignment(
#                 vertical="center",
#                 wrap_text=True
#             )

#     # ======================================================
#     # API TEST RESULTS SHEET
#     # ======================================================

#     api_sheet = workbook.create_sheet("API Test Results")

#     api_sheet.append([
#         "Sr. No.",
#         "API Name",
#         "Staff Number",
#         "Test Case ID",
#         "Test Type",
#         "Priority",
#         "Test Scenario",
#         "Method",
#         "Endpoint",
#         "Test Data / Payload",
#         "Status Code",
#         "Expected Result",
#         "Actual Result",
#         "Response Message",
#         "Speed",
#         "Pass/Fail",
#         "Response Time (ms)",
#         "Error Details",
#         "Execution Date & Time"
#     ])

#     for index, result in enumerate(api_results, start=1):

#         api_sheet.append([
#             index,
#             result.get("API Name", ""),
#             result.get("Staff Number", ""),
#             result.get("Test Case ID", ""),
#             result.get("Test Type", ""),
#             result.get("Priority", ""),
#             result.get("Test Scenario", ""),
#             result.get("Method", ""),
#             result.get("Endpoint", ""),
#             str(result.get("Test Data / Payload", "")),
#             result.get("Status Code", ""),
#             result.get("Expected Result", ""),
#             result.get("Actual Result", ""),
#             result.get("Response Message", ""),
#             result.get("Speed", ""),
#             result.get("Status", ""),
#             result.get("Response Time (ms)", ""),
#             result.get("Error", ""),
#             result.get("Execution Date & Time", "")
#         ])

#     # ======================================================
#     # API RESULTS FORMATTING
#     # ======================================================

#     for cell in api_sheet[1]:
#         cell.font = Font(bold=True)
#         cell.alignment = Alignment(
#             horizontal="center",
#             vertical="center",
#             wrap_text=True
#         )

#     for row in api_sheet.iter_rows(min_row=2):
#         for cell in row:
#             cell.alignment = Alignment(
#                 vertical="top",
#                 wrap_text=True
#             )

#     api_sheet.freeze_panes = "A2"
#     api_sheet.auto_filter.ref = api_sheet.dimensions
#     api_sheet.row_dimensions[1].height = 35

#     summary.freeze_panes = "A3"

#     # ======================================================
#     # AUTO COLUMN WIDTH
#     # ======================================================

#     set_column_widths(workbook)

#     # ======================================================
#     # SAVE API REPORT
#     # ======================================================

#     workbook.save(REPORT_FILE)

#     print("\nAPI REPORT GENERATED SUCCESSFULLY")
#     print("Total API Tests:", total)
#     print("Passed:", passed)
#     print("Failed:", failed)
#     print("Report Status:", report_status)
#     print("Report Path:", os.path.abspath(REPORT_FILE))


# # ==========================================================
# # CREATE REPORT
# # ==========================================================

# def create_report(results, api_results=None, api_only=False):

#     # API-only report: keep Selenium report unchanged
#     if api_only:
#         create_api_report(api_results or [])
#         return

#     workbook = Workbook()

#     # ======================================================
#     # SUMMARY SHEET
#     # ======================================================

#     summary = workbook.active
#     summary.title = "Summary"

#     # ======================================================
#     # GET MODULES DYNAMICALLY
#     # ======================================================

#     modules = []

#     for result in results:

#         module = result.get(
#             "Module",
#             result.get("Test Case", "Unknown")
#         )

#         if module not in modules:
#             modules.append(module)

#     # ======================================================
#     # SUMMARY TITLE
#     # ======================================================

#     summary["A1"] = "Test Execution Summary"
#     summary["A1"].font = Font(bold=True, size=14)
#     summary.merge_cells("A1:I1")
#     summary["A1"].alignment = Alignment(horizontal="center")

#     # ======================================================
#     # SUMMARY HEADER
#     # ======================================================

#     summary.append([])

#     summary.append([
#         "Module",
#         "Total Test Cases",
#         "Passed",
#         "Failed",
#         "First Attempt Failures",
#         "Retry Attempts",
#         "Passed After Retry",
#         "Total Attempts",
#         "Pass %"
#     ])

#     # ======================================================
#     # MODULE-WISE SUMMARY
#     # ======================================================

#     for module in modules:

#         module_results = [
#             result for result in results
#             if result.get(
#                 "Module",
#                 result.get("Test Case", "Unknown")
#             ) == module
#         ]

#         # Unique test data
#         test_cases = []

#         for result in module_results:

#             test_data = result.get("Test Data", "")

#             if test_data not in test_cases:
#                 test_cases.append(test_data)

#         total = len(test_cases)

#         # Final status
#         final_status = {}

#         for result in module_results:

#             test_data = result.get("Test Data", "")
#             status = result.get("Final Status", "")

#             if status in ("PASS", "FAIL"):
#                 final_status[test_data] = status

#         passed = sum(
#             1 for status in final_status.values()
#             if status == "PASS"
#         )

#         failed = sum(
#             1 for status in final_status.values()
#             if status == "FAIL"
#         )

#         # First attempt failures
#         first_attempt_failures = sum(
#             1 for result in module_results
#             if result.get("Attempt") == 1
#             and result.get("Status") == "FAIL"
#         )

#         # Retry attempts
#         retry_attempts = sum(
#             1 for result in module_results
#             if result.get("Final Status") == "RETRY"
#         )

#         # Passed after retry
#         passed_after_retry = 0

#         for test_data in test_cases:

#             if final_status.get(test_data) == "PASS":

#                 retry_pass = any(
#                     result.get("Test Data") == test_data
#                     and result.get("Attempt", 1) > 1
#                     and result.get("Status") == "PASS"
#                     for result in module_results
#                 )

#                 if retry_pass:
#                     passed_after_retry += 1

#         # Total attempts
#         total_attempts = len(module_results)

#         # Pass percentage
#         pass_percentage = (
#             (passed / total) * 100
#             if total > 0 else 0
#         )

#         summary.append([
#             module,
#             total,
#             passed,
#             failed,
#             first_attempt_failures,
#             retry_attempts,
#             passed_after_retry,
#             total_attempts,
#             f"{pass_percentage:.2f}%"
#         ])

#     # ======================================================
#     # SUMMARY FORMATTING
#     # ======================================================

#     for cell in summary[3]:
#         cell.font = Font(bold=True)
#         cell.alignment = Alignment(
#             horizontal="center",
#             vertical="center",
#             wrap_text=True
#         )

#     for row in summary.iter_rows(min_row=4):
#         for cell in row:
#             cell.alignment = Alignment(
#                 horizontal="center",
#                 vertical="center",
#                 wrap_text=True
#             )

#     # ======================================================
#     # EXECUTION INFORMATION
#     # ======================================================

#     execution_row = summary.max_row + 2

#     summary.cell(execution_row, 1).value = "Execution Information"
#     summary.cell(execution_row, 1).font = Font(bold=True, size=12)

#     summary.cell(execution_row + 1, 1).value = "Execution Date"
#     summary.cell(execution_row + 1, 2).value = (
#         datetime.now().strftime("%d-%m-%Y")
#     )

#     summary.cell(execution_row + 2, 1).value = "Execution Time"
#     summary.cell(execution_row + 2, 2).value = (
#         datetime.now().strftime("%H:%M:%S")
#     )

#     # ======================================================
#     # OVERALL REPORT STATUS
#     # ======================================================

#     all_final_statuses = [
#         result for result in results
#         if result.get("Final Status") in ("PASS", "FAIL")
#     ]

#     overall_failed = any(
#         result.get("Final Status") == "FAIL"
#         for result in all_final_statuses
#     )

#     if not results:
#         report_status = "NO DATA"
#     elif overall_failed:
#         report_status = "FAILED"
#     else:
#         report_status = "PASSED"

#     summary.cell(execution_row + 3, 1).value = "Overall Status"
#     summary.cell(execution_row + 3, 2).value = report_status

#     summary.cell(execution_row + 4, 1).value = "Report File"
#     summary.cell(execution_row + 4, 2).value = os.path.basename(REPORT_FILE)

#     for row in range(execution_row + 1, execution_row + 5):
#         summary.cell(row, 1).font = Font(bold=True)

   
#     # ======================================================
#     # TEST RESULTS SHEET - WEB AUTOMATION
#     # ======================================================

#     details = workbook.create_sheet("Test Results")

#     details.append([
#         "Sr. No.",
#         "Module",
#         "Test Case",
#         "Test Data",
#         "Attempt",
#         "Status",
#         "Expected Result",
#         "Actual Result",
#         "Error Details",
#         "Screenshot Path",
#         "Execution Time (seconds)",
#         "Final Status"
#     ])

#     for index, result in enumerate(results, start=1):

#         module = result.get(
#             "Module",
#             result.get("Test Case", "Unknown")
#         )

#         details.append([
#             index,
#             module,
#             result.get("Test Case", ""),
#             result.get("Test Data", ""),
#             result.get("Attempt", ""),
#             result.get("Status", ""),
#             result.get("Expected Result", ""),
#             result.get("Actual Result", ""),
#             result.get("Error Details", result.get("Error", "")),
#             result.get("Screenshot Path", ""),
#             result.get("Execution Time (seconds)", ""),
#             result.get("Final Status", "")
#         ])

#     # Format headers
#     for cell in details[1]:
#         cell.font = Font(bold=True)
#         cell.alignment = Alignment(
#             horizontal="center",
#             vertical="center",
#             wrap_text=True
#         )

#     # Format result cells
#     for row in details.iter_rows(min_row=2):
#         for cell in row:
#             cell.alignment = Alignment(
#                 vertical="top",
#                 wrap_text=True
#             )

#     details.row_dimensions[1].height = 35

#     if details.max_row > 1:
#         details.auto_filter.ref = details.dimensions

#     details.freeze_panes = "A2"


#     # ======================================================
#     # AUTO COLUMN WIDTH
#     # ======================================================

#     set_column_widths(workbook)

#     # ======================================================
#     # SAVE REPORT
#     # ======================================================

#     workbook.save(REPORT_FILE)

#     # ======================================================
#     # CONSOLE OUTPUT
#     # ======================================================

#     print("\n====================================")
#     print("REPORT GENERATED SUCCESSFULLY")
#     print("====================================")
#     print(f"Total Results : {len(results)}")
#     print(f"Report Status : {report_status}")
#     print(f"Report Path   : {os.path.abspath(REPORT_FILE)}")
#     print("====================================")



from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import os


# ==========================================================
# DOWNLOAD FOLDER
# ==========================================================

DOWNLOAD_FOLDER = os.path.join(
    os.path.expanduser("~"),
    "Downloads"
)

os.makedirs(DOWNLOAD_FOLDER, exist_ok=True)


# ==========================================================
# REPORT FILE
# ==========================================================

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

REPORT_FILE = os.path.join(
    DOWNLOAD_FOLDER,
    f"test_execution_report_{timestamp}.xlsx"
)


# ==========================================================
# COMMON: AUTO COLUMN WIDTH
# ==========================================================

def set_column_widths(workbook):

    for sheet in workbook.worksheets:
        for column in sheet.columns:

            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            sheet.column_dimensions[column_letter].width = min(
                max_length + 3,
                50
            )


# ==========================================================
# API REPORT
# EXISTING API REPORTING LOGIC PRESERVED
# ==========================================================

def create_api_report(api_results):

    workbook = Workbook()

    # ======================================================
    # API SUMMARY SHEET
    # ======================================================

    summary = workbook.active
    summary.title = "Summary"

    total = len(api_results)

    passed = sum(
        1 for result in api_results
        if result.get("Status") == "PASS"
    )

    failed = sum(
        1 for result in api_results
        if result.get("Status") == "FAIL"
    )

    pass_percentage = (
        (passed / total) * 100
        if total > 0 else 0
    )

    summary["A1"] = "API Test Execution Summary"
    summary["A1"].font = Font(bold=True, size=14)
    summary.merge_cells("A1:B1")
    summary["A1"].alignment = Alignment(horizontal="center")

    summary.append(["Metric", "Value"])
    summary.append(["Total API Tests", total])
    summary.append(["Passed", passed])
    summary.append(["Failed", failed])
    summary.append(["Pass Percentage", f"{pass_percentage:.2f}%"])
    summary.append([
        "Execution Date",
        datetime.now().strftime("%d-%m-%Y")
    ])
    summary.append([
        "Execution Time",
        datetime.now().strftime("%H:%M:%S")
    ])

    if total == 0:
        report_status = "NO DATA"
    elif failed > 0:
        report_status = "FAILED"
    else:
        report_status = "PASSED"

    summary.append(["Overall Status", report_status])
    summary.append(["Report File", os.path.basename(REPORT_FILE)])

    for cell in summary[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in summary.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

    # ======================================================
    # API TEST RESULTS SHEET
    # ======================================================

    api_sheet = workbook.create_sheet("API Test Results")

    api_sheet.append([
        "Sr. No.",
        "API Name",
        "Staff Number",
        "Test Case ID",
        "Test Type",
        "Priority",
        "Test Scenario",
        "Method",
        "Endpoint",
        "Test Data / Payload",
        "Status Code",
        "Expected Result",
        "Actual Result",
        "Response Message",
        "Speed",
        "Pass/Fail",
        "Response Time (ms)",
        "Error Details",
        "Execution Date & Time"
    ])

    for index, result in enumerate(api_results, start=1):

        api_sheet.append([
            index,
            result.get("API Name", ""),
            result.get("Staff Number", ""),
            result.get("Test Case ID", ""),
            result.get("Test Type", ""),
            result.get("Priority", ""),
            result.get("Test Scenario", ""),
            result.get("Method", ""),
            result.get("Endpoint", ""),
            str(result.get("Test Data / Payload", "")),
            result.get("Status Code", ""),
            result.get("Expected Result", ""),
            result.get("Actual Result", ""),
            result.get("Response Message", ""),
            result.get("Speed", ""),
            result.get("Status", ""),
            result.get("Response Time (ms)", ""),
            result.get("Error", ""),
            result.get("Execution Date & Time", "")
        ])

    # ======================================================
    # API RESULTS FORMATTING
    # ======================================================

    for cell in api_sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in api_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    api_sheet.freeze_panes = "A2"
    api_sheet.auto_filter.ref = api_sheet.dimensions
    api_sheet.row_dimensions[1].height = 35

    summary.freeze_panes = "A3"

    # ======================================================
    # AUTO COLUMN WIDTH
    # ======================================================

    set_column_widths(workbook)

    # ======================================================
    # SAVE API REPORT
    # ======================================================

    workbook.save(REPORT_FILE)

    print("\nAPI REPORT GENERATED SUCCESSFULLY")
    print("Total API Tests:", total)
    print("Passed:", passed)
    print("Failed:", failed)
    print("Report Status:", report_status)
    print("Report Path:", os.path.abspath(REPORT_FILE))


# ==========================================================
# CREATE WEB AUTOMATION REPORT
# ==========================================================

def create_report(results, api_results=None, api_only=False):

    # Keep API-only reporting behavior unchanged
    if api_only:
        create_api_report(api_results or [])
        return

    workbook = Workbook()

    # ======================================================
    # PREPARE TEST IDENTIFIERS
    # ======================================================

    def get_module(result):
        return result.get(
            "Module",
            result.get("Test Case", "Unknown")
        )

    def get_test_key(result):
        """
        Identify one logical test using Module, Test Case
        and Test Data. Different test data stays separate.
        """

        return (
            get_module(result),
            result.get("Test Case", ""),
            str(result.get("Test Data", ""))
        )

    def get_status(result):
        return str(result.get("Status", "")).upper()

    # Group all attempts belonging to the same logical test
    grouped_tests = {}

    for result in results:

        key = get_test_key(result)

        if key not in grouped_tests:
            grouped_tests[key] = []

        grouped_tests[key].append(result)

    # Sort attempts in execution order
    for attempts in grouped_tests.values():
        attempts.sort(
            key=lambda item: int(item.get("Attempt", 1) or 1)
        )

    # Determine the final status for each logical test
    final_statuses = {}

    for key, attempts in grouped_tests.items():

        # Prefer an explicit final status from the input
        explicit_statuses = [
            str(item.get("Final Status", "")).upper()
            for item in attempts
            if str(item.get("Final Status", "")).upper()
            in ("PASS", "FAIL")
        ]

        if explicit_statuses:
            final_status = explicit_statuses[-1]
        else:
            # Otherwise use the last execution attempt
            final_status = get_status(attempts[-1])

        final_statuses[key] = final_status

    # ======================================================
    # SUMMARY SHEET
    # ======================================================

    summary = workbook.active
    summary.title = "Summary"

    modules = list(dict.fromkeys(
        key[0] for key in grouped_tests
    ))

    summary["A1"] = "Test Execution Summary"
    summary["A1"].font = Font(bold=True, size=14)
    summary.merge_cells("A1:I1")
    summary["A1"].alignment = Alignment(horizontal="center")

    summary.append([])

    summary.append([
        "Module",
        "Total Test Cases",
        "Passed",
        "Failed",
        "First Attempt Failures",
        "Retry Attempts",
        "Passed After Retry",
        "Total Attempts",
        "Pass %"
    ])

    # ======================================================
    # MODULE-WISE SUMMARY
    # ======================================================

    for module in modules:

        module_keys = [
            key for key in grouped_tests
            if key[0] == module
        ]

        total = len(module_keys)

        passed = sum(
            1 for key in module_keys
            if final_statuses.get(key) == "PASS"
        )

        failed = sum(
            1 for key in module_keys
            if final_statuses.get(key) == "FAIL"
        )

        first_attempt_failures = 0
        retry_attempts = 0
        passed_after_retry = 0
        total_attempts = 0

        for key in module_keys:

            attempts = grouped_tests[key]
            total_attempts += len(attempts)

            # First execution attempt
            first_attempt = next(
                (
                    item for item in attempts
                    if int(item.get("Attempt", 1) or 1) == 1
                ),
                attempts[0]
            )

            if get_status(first_attempt) == "FAIL":
                first_attempt_failures += 1

            # Count all attempts after the first attempt
            retry_attempts += sum(
                1 for item in attempts
                if int(item.get("Attempt", 1) or 1) > 1
            )

            # Test failed initially but passed on a later attempt
            if (
                get_status(first_attempt) == "FAIL"
                and final_statuses.get(key) == "PASS"
            ):
                passed_after_retry += 1

        pass_percentage = (
            (passed / total) * 100
            if total > 0 else 0
        )

        summary.append([
            module,
            total,
            passed,
            failed,
            first_attempt_failures,
            retry_attempts,
            passed_after_retry,
            total_attempts,
            f"{pass_percentage:.2f}%"
        ])

    # ======================================================
    # SUMMARY FORMATTING
    # ======================================================

    for cell in summary[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in summary.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

    # ======================================================
    # EXECUTION INFORMATION
    # ======================================================

    execution_row = summary.max_row + 2

    summary.cell(execution_row, 1).value = "Execution Information"
    summary.cell(execution_row, 1).font = Font(
        bold=True,
        size=12
    )

    summary.cell(execution_row + 1, 1).value = "Execution Date"
    summary.cell(execution_row + 1, 2).value = (
        datetime.now().strftime("%d-%m-%Y")
    )

    summary.cell(execution_row + 2, 1).value = "Execution Time"
    summary.cell(execution_row + 2, 2).value = (
        datetime.now().strftime("%H:%M:%S")
    )

    # ======================================================
    # OVERALL REPORT STATUS
    # ======================================================

    if not grouped_tests:
        report_status = "NO DATA"

    elif any(
        status == "FAIL"
        for status in final_statuses.values()
    ):
        report_status = "FAILED"

    else:
        report_status = "PASSED"

    summary.cell(execution_row + 3, 1).value = "Overall Status"
    summary.cell(execution_row + 3, 2).value = report_status

    summary.cell(execution_row + 4, 1).value = "Report File"
    summary.cell(execution_row + 4, 2).value = (
        os.path.basename(REPORT_FILE)
    )

    for row in range(execution_row + 1, execution_row + 5):
        summary.cell(row, 1).font = Font(bold=True)

    summary.freeze_panes = "A4"

    # ======================================================
    # TEST RESULTS SHEET - WEB AUTOMATION
    # ======================================================

    details = workbook.create_sheet("Test Results")

    details.append([
        "Sr. No.",
        "Module",
        "Test Case",
        "Test Data",
        "Attempt",
        "Status",
        "Expected Result",
        "Actual Result",
        "Error Details",
        "Screenshot Path",
        "Execution Time (seconds)",
        "Final Status"
    ])

    index = 1

    # Write each attempt as a separate row
    for key, attempts in grouped_tests.items():

        final_status = final_statuses.get(key, "")

        for result in attempts:

            module = get_module(result)

            details.append([
                index,
                module,
                result.get("Test Case", ""),
                result.get("Test Data", ""),
                result.get("Attempt", 1),
                result.get("Status", ""),
                result.get("Expected Result", ""),
                result.get("Actual Result", ""),
                result.get(
                    "Error Details",
                    result.get("Error", "")
                ),
                result.get("Screenshot Path", ""),
                result.get("Execution Time (seconds)", ""),
                final_status
            ])

            index += 1

    # ======================================================
    # TEST RESULTS FORMATTING
    # ======================================================

    for cell in details[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    details.row_dimensions[1].height = 35

    if details.max_row > 1:
        details.auto_filter.ref = details.dimensions

    details.freeze_panes = "A2"

    # ======================================================
    # AUTO COLUMN WIDTH
    # ======================================================

    set_column_widths(workbook)

    # ======================================================
    # SAVE REPORT
    # ======================================================

    workbook.save(REPORT_FILE)

    # ======================================================
    # CONSOLE OUTPUT
    # ======================================================

    total_tests = len(grouped_tests)

    total_passed = sum(
        1 for status in final_statuses.values()
        if status == "PASS"
    )

    total_failed = sum(
        1 for status in final_statuses.values()
        if status == "FAIL"
    )

    print("\n====================================")
    print("WEB REPORT GENERATED SUCCESSFULLY")
    print("====================================")
    print(f"Total Test Cases : {total_tests}")
    print(f"Passed           : {total_passed}")
    print(f"Failed           : {total_failed}")
    print(f"Total Attempts   : {len(results)}")
    print(f"Report Status    : {report_status}")
    print(f"Report Path      : {os.path.abspath(REPORT_FILE)}")
    print("====================================")
