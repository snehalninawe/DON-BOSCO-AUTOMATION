
from openpyxl import load_workbook
import os

EXCEL_FILE = "test_data.xlsx"
API_EXCEL_FILE = "Api_Test_data.xlsx"


def read_excel(file_name, sheet_name):
    file_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        file_name
    )

    workbook = load_workbook(file_path)

    try:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        all_data = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in row):
                continue

            data = {
                header: row[index]
                for index, header in enumerate(headers)
                if header is not None
            }

            all_data.append(data)

        return all_data

    finally:
        workbook.close()


def get_all_data(sheet_name):
    return read_excel(EXCEL_FILE, sheet_name)


def get_data(sheet_name):
    data = get_all_data(sheet_name)

    if not data:
        raise ValueError(f"No test data found in sheet: {sheet_name}")

    # Return the first row as a dictionary
    return data[0]


def get_api_data(sheet_name):
    return read_excel(API_EXCEL_FILE, sheet_name)



def read_excel(file_name, sheet_name):
    file_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        file_name
    )

    workbook = load_workbook(file_path)

    try:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        all_data = []

        for row in sheet.iter_rows(min_row=2, values_only=True):

            # Skip completely empty rows
            if all(value is None for value in row):
                continue

            data = {}

            for column, header in enumerate(headers):
                data[header] = row[column]

            all_data.append(data)

        return all_data

    finally:
        workbook.close()
